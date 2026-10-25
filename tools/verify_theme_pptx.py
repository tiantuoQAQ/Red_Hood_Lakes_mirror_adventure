# -*- coding: utf-8 -*-
"""校验生成的 PPT：页数、文本、图片数量、越界检查。"""
from pptx import Presentation
from pptx.util import Emu

path = "docs/红之镜-作品与主题说明.pptx"
prs = Presentation(path)
W, H = prs.slide_width, prs.slide_height
print("slide size:", W / 914400, "x", H / 914400, "in | slides:", len(prs.slides))

problems = 0
for si, slide in enumerate(prs.slides, 1):
    print("\n--- Slide", si, "---")
    for shape in slide.shapes:
        left, top = shape.left or 0, shape.top or 0
        w, h = shape.width or 0, shape.height or 0
        if left + w > W + 9144 or top + h > H + 9144 or left < -9144 or top < -9144:
            problems += 1
            print("  !! OUT OF BOUNDS:", shape.shape_type, left / 914400, top / 914400, w / 914400, h / 914400)
        if shape.has_text_frame and shape.text_frame.text.strip():
            t = shape.text_frame.text.strip().replace("\n", " / ")
            print("  [txt]", t[:70], "| size:", round(shape.width / 914400, 2), "x", round(shape.height / 914400, 2))
        if shape.shape_type is not None and shape.shape_type == 13:
            print("  [pic]", round(shape.width / 914400, 2), "x", round(shape.height / 914400, 2))
print("\nproblems:", problems)
