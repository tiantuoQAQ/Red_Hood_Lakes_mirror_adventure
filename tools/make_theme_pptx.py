# -*- coding: utf-8 -*-
"""生成《红之镜 DEMO》主题说明 PPT（2 页）：作品名 + 与「红之镜」主题的关联。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

RED = RGBColor(0xFF, 0x00, 0x00)
DARK = RGBColor(0x0A, 0x0A, 0x0C)
PANEL = RGBColor(0x14, 0x14, 0x18)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xAA, 0xAA, 0xB0)
DGRAY = RGBColor(0x66, 0x66, 0x6E)
CYAN = RGBColor(0x00, 0xBA, 0xFF)
GREEN = RGBColor(0x82, 0xFF, 0x00)
YAHEI = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bg_rect(slide, color=DARK):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def text(slide, x, y, w, h, content, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         font=YAHEI, line_spacing=1.15, space_after=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box


def line(slide, x, y, w, color=RED, thickness=2.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(thickness))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def panel(slide, x, y, w, h, color=PANEL, border=DGRAY):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = border
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


# ============ 第 1 页：作品名 ============
s1 = prs.slides.add_slide(prs.slide_layouts[6])
bg_rect(s1)

# 左侧竖条装饰
bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.5), Inches(0.09), Inches(3.0))
bar.fill.solid()
bar.fill.fore_color.rgb = RED
bar.line.fill.background()
bar.shadow.inherit = False

text(s1, 0.85, 1.45, 6.2, 0.4, "CODEJAM 2026  ·  SUDA CEA  ·  DEMO", 14, GRAY)
text(s1, 0.85, 1.9, 6.4, 1.5, "红之镜", 92, RED, bold=True, line_spacing=0.95)
text(s1, 0.9, 3.45, 6.2, 0.5, "MIRROR  OF  RED", 24, WHITE, line_spacing=1.0)
line(s1, 0.9, 3.95, 3.4, RED, 2.5)
text(s1, 0.9, 4.15, 6.0, 1.2,
     ["弹幕闪避 × 平台跳跃  ·  横版 ACT 原型",
      "一人一镜，穿过「红」的倒影，去砍断写故事的人。"],
     15, GRAY, line_spacing=1.3)

# 右侧：标题画（游戏真实背景）
pic = s1.shapes.add_picture("art/bg/bg_mirrorworld.png", Inches(7.3), Inches(0.9), width=Inches(5.5))
pic2 = s1.shapes.add_picture("art/logo/hongzhijing_logo.png", Inches(7.75), Inches(2.55), width=Inches(4.6))

text(s1, 7.55, 5.05, 5.2, 1.6,
     ["主题关键词",
      "红 —— 血、罪、警示，也是蕾克唯一的颜色",
      "镜 —— 倒影、翻版、另一个自己",
      "红之镜 —— 照见「相反的她们」的窗"],
     13, GRAY, line_spacing=1.25)

text(s1, 0.85, 6.5, 11.8, 0.4, "Godot 4.7.2  ·  1280×720  ·  键盘 / 鼠标  ·  单机 Demo", 12, DGRAY)

# ============ 第 2 页：与主题的关联 ============
s2 = prs.slides.add_slide(prs.slide_layouts[6])
bg_rect(s2)

text(s2, 0.6, 0.42, 8.0, 0.6, "为什么是「红之镜」", 34, RED, bold=True, line_spacing=1.0)
text(s2, 0.62, 1.02, 8.6, 0.4, "镜中世界关着一整套「主世界人物的相反面」，而蕾克手里握着穿越镜面的权能", 14, GRAY)
line(s2, 0.62, 1.5, 12.1, DGRAY, 1.0)

# 左：核心设定
panel(s2, 0.6, 1.75, 5.55, 4.9)
text(s2, 0.85, 1.95, 5.1, 0.4, "核心设定", 17, WHITE, bold=True)
line(s2, 0.85, 2.32, 4.9, RED, 2.0)
text(s2, 0.85, 2.5, 5.05, 3.9,
     ["· 蕾克与「镜」本就脱不开关系",
      "  原作中她就与镜世界纠缠不清，本作保留了她「穿越镜中世界」的权能。",
      "· 镜中世界是主世界的翻版",
      "  镜里出现的「她们」，都是主世界线人物的相反面——",
      "  同一个人的另一条可能。",
      "· 但玛丽苏不是翻版",
      "  她不属于任何人的倒影，她是「写故事的人」：凌驾于镜像之上的作者，也是那句「恭喜你，走到这里」的施舍者。"],
     13.5, GRAY, line_spacing=1.3)

# 右：七宗罪谱系
text(s2, 6.55, 1.75, 6.0, 0.4, "镜中世界的「罪」谱系", 17, WHITE, bold=True)
line(s2, 6.55, 2.12, 5.9, RED, 2.0)

rows = [
    ("贪婪 · 艾露玛", "主世界：艾露玛  ·  相反面：永不满足的占有欲", GRAY),
    ("傲慢 · 贞德", "主世界：贞德  ·  相反面：不容置疑的自我神化", GRAY),
    ("惰怠 · 维多利亚", "主世界：维多利亚  ·  相反面：放弃一切的停滞", GRAY),
    ("暴虐 · 猩红魔姬", "主世界：蕾克  ·  相反面：把恨当力量的自己  ← Demo 已实装", RED),
    ("罪人 · 格林", "主世界：格林  ·  相反面：背负罪名的那条路", GRAY),
    ("作者 · 玛丽苏", "唯一不是翻版：镜中世界的书写者  ← Demo 最终战", GREEN),
]
y = 2.32
for name, desc, color in rows:
    text(s2, 6.6, y, 5.6, 0.32, name, 14.5, color, bold=True, line_spacing=1.0)
    text(s2, 6.6, y + 0.30, 5.9, 0.3, desc, 11.5, DGRAY, line_spacing=1.0)
    y += 0.72

# 底部：Demo 落点
panel(s2, 0.6, 6.72, 12.13, 0.66, PANEL, RED)
text(s2, 0.85, 6.82, 11.7, 0.5,
     ["Demo 落点：以「暴虐 · 猩红魔姬」与「作者 · 玛丽苏」两场镜像之战，验证「红之镜」的核心命题——",
      "你要砍断的，是别人的故事，也是自己的倒影。"],
     12, WHITE, line_spacing=1.05, space_after=1)

out = "docs/红之镜-作品与主题说明.pptx"
prs.save(out)
print("saved", out, "slides =", len(prs.slides.__iter__.__self__._sldIdLst))
